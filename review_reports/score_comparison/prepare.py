"""Join the pathologists' readings with both AI score files, one row per case and marker."""
import csv, json, pathlib, re, statistics as st, sys
import openpyxl

# The workspace root, from this file's own place - review_reports/<review>/<script>.py
# - rather than one machine's path (P-20).
REPO = pathlib.Path(__file__).resolve().parents[2]
sys.path.append(str(REPO))  # data_versions.py lives at the workspace root
import data_versions  # noqa: E402
#: Intermediate files (tables, images, charts) live under v<N>_data/data/, not beside the code.
HERE = data_versions.data_root() / "review_reports" / "score_comparison"
HERE.mkdir(parents=True, exist_ok=True)
MARKERS = ["A", "F", "R", "U", "W"]

wb = openpyxl.load_workbook(data_versions.SHARED_DATA_ROOT / "oncostem_docs/client/OncoStem/6Slide Reports2.xlsx", data_only=True)
ws = wb.worksheets[0]
header = [c.value for c in ws[1]]
readers = {}  # (case, marker) -> {"readers": {id: (pct, int)}, "avg": (pct, int)}
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row[0]:
        continue
    who, rest = row[0].split("_", 1)
    case = "CAN_" + re.search(r"CAN[_/]?(\d{5})", rest).group(1)
    values = dict(zip(header[1:], row[1:]))
    for m in MARKERS:
        pk = next(k for k in values if k.startswith(f"{m}%"))
        ik = next(k for k in values if k.startswith(f"{m}I"))
        slot = readers.setdefault((case, m), {"readers": {}, "avg": None})
        if who == "AVG":
            slot["avg"] = (float(values[pk]), float(values[ik]))
        else:
            slot["readers"][who] = (float(values[pk]), float(values[ik]))

def load(name):
    out = {}
    with open(data_versions.results_root() / name, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            out[(r["case_id"], r["marker"])] = r
    return out
pixel, tile = load("oncostem_ai_scores.csv"), load("oncostem_ai_scores_tiles.csv")

rows = []
for (case, m), slot in sorted(readers.items()):
    pv = [v[0] for v in slot["readers"].values()]
    iv = [v[1] for v in slot["readers"].values()]
    p, t = pixel.get((case, m)), tile.get((case, m))
    rows.append({
        "case": case, "marker": m, "markerName": (p or t)["marker_name"],
        "readers": slot["readers"],
        "readerPctMean": round(st.mean(pv), 2), "readerPctMin": min(pv), "readerPctMax": max(pv),
        "readerIntMean": round(st.mean(iv), 3), "readerIntMin": min(iv), "readerIntMax": max(iv),
        "avgSheet": slot["avg"],
        "pixel": None if not p or p["state"] != "scored" else {"pct": float(p["percent"]), "int": float(p["intensity"]), "label": p["intensity_label"], "cells": int(p["cells"] or 0)},
        "tile": None if not t or t["state"] != "scored" else {"pct": float(t["percent"]), "int": float(t["intensity"]), "label": t["intensity_label"], "cells": int(t["cells"] or 0)},
        "provisional": (p or t)["cuts_provisional"],
    })
(HERE / "rows.json").write_text(json.dumps(rows, indent=1), encoding="utf-8")

# agreement summary
def summary(key, field):
    diffs, within = [], 0
    for r in rows:
        s = r[key]
        if not s: continue
        ref = r["readerPctMean"] if field == "pct" else r["readerIntMean"]
        lo, hi = (r["readerPctMin"], r["readerPctMax"]) if field == "pct" else (r["readerIntMin"], r["readerIntMax"])
        diffs.append(s[field] - ref); within += lo <= s[field] <= hi
    return len(diffs), round(st.mean(abs(d) for d in diffs), 2), round(st.mean(diffs), 2), within
for key in ("pixel", "tile"):
    print(key, "pct n/MAE/bias/within", summary(key, "pct"), "int", summary(key, "int"))
# reader-vs-reader spread for scale
spread = [st.mean(abs(v[0] - r["readerPctMean"]) for v in r["readers"].values()) for r in rows]
print("readers' own mean abs deviation from their mean, pct:", round(st.mean(spread), 2))
print("avg sheet vs computed mean mismatches:", [(r["case"], r["marker"], r["avgSheet"], r["readerPctMean"], r["readerIntMean"]) for r in rows if abs(r["avgSheet"][0] - r["readerPctMean"]) > 0.01 or abs(r["avgSheet"][1] - r["readerIntMean"]) > 0.02])
for m in MARKERS:
    sub = [r for r in rows if r["marker"] == m]
    print(m, sub[0]["markerName"], "pixel MAE", round(st.mean(abs(r["pixel"]["pct"] - r["readerPctMean"]) for r in sub if r["pixel"]), 1), "tile MAE", round(st.mean(abs(r["tile"]["pct"] - r["readerPctMean"]) for r in sub if r["tile"]), 1))
print(sum(1 for r in rows if not r["pixel"]), "missing pixel;", sum(1 for r in rows if not r["tile"]), "missing tile")
