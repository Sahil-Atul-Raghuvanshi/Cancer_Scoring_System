"""One case's results, as the workbook somebody outside this system opens.

Seven sheets, and the order is the argument:

    Scores        the ten numbers, and nothing else on the sheet that could be
                  mistaken for them
    Agreement     our numbers against the pathologists', per marker, with the
                  four individual reads and the +/-10 human spread. Second,
                  because it is the only sheet that says whether the first one
                  is any good
    Workings      how each pair was arrived at, line by line, with the numbers
                  substituted in - so the Scores sheet has no hidden arithmetic
                  behind it
    Caveats       what has to be known before the ten numbers are used
    Reference     H-score, Allred, the ASCO/CAP call - the field's standard
                  vocabulary, on their own sheet precisely so they cannot be
                  mistaken for the deliverable
    Regions       each invasive region's own answer, and the two ways of
                  combining them
    Run           which slides, which steps, how long, and what failed

**The deliverable sheet is first and holds two columns per marker.** OncoStem's
own sheet has ten columns - five percent-and-intensity pairs - and no combined
score anywhere. Putting an H-score next to the pair on the same sheet is how a
reader starts treating it as the answer, so it is on a different one, under a
heading that says what it is for.

**A marker that did not finish gets a row saying so, not a blank and not a
zero.** A missing measurement and a measured zero are different facts, and a
spreadsheet is exactly the place where they stop looking different.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

from app import panel
from app.core.logging import get_logger
from app.scoring import cuts as cut_points
from app.services.case_score_service import case_score_service
from app.services.per_cell_service import per_cell_service
from app.services.score_service import ScoreError, score_service

logger = get_logger(__name__)

_HEADER_FILL = "FFE8EEF7"
_DELIVERABLE_FILL = "FFDDF0DD"
_WARN_FILL = "FFFDE8E8"


def _style_header(sheet, row: int = 1) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor=_HEADER_FILL)
    for cell in sheet[row]:
        if cell.value is None:
            continue
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _autosize(sheet, *, maximum: int = 70) -> None:
    from openpyxl.utils import get_column_letter

    for index, column in enumerate(sheet.columns, start=1):
        width = max(
            (len(str(cell.value)) for cell in column if cell.value is not None),
            default=10,
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(
            maximum, max(10, width + 2)
        )


def _wrap(sheet, column: str, width: int = 90) -> None:
    from openpyxl.styles import Alignment

    sheet.column_dimensions[column].width = width
    for cell in sheet[column]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_case_workbook(
    case_id: str,
    destination: Path,
    *,
    runs: list[dict] | None = None,
    case_path: str = "",
) -> Path:
    """Write `destination` from whatever of `case_id` has been scored.

    `runs` is the batch runner's own per-marker log. It is optional: the
    workbook is built from the reports on disk, so it can be regenerated at any
    time without re-running anything.
    """
    from openpyxl import Workbook

    book = Workbook()
    grid = case_score_service.report(case_id)
    pairs = case_score_service.pairs(case_id)

    scores: dict[str, object] = {}
    measurements: dict[str, object] = {}
    for letter, (he, ihc) in pairs.items():
        try:
            scores[letter] = score_service.report(he, ihc).score
            measurements[letter] = per_cell_service.report(he, ihc)
        except (ScoreError, Exception):  # noqa: BLE001 - a missing marker is a row, not a crash
            continue

    _scores_sheet(book, grid, scores, case_id, case_path)
    _agreement_sheet(book, scores, case_id)
    _workings_sheet(book, scores)
    _caveats_sheet(book, scores)
    _reference_sheet(book, scores)
    _regions_sheet(book, scores, measurements)
    _run_sheet(book, runs or [], case_id, case_path)

    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)
    logger.info("wrote %s", destination)
    return destination


# --- sheet 1: the deliverable ---------------------------------------------


def _scores_sheet(book, grid, scores: dict, case_id: str, case_path: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet = book.active
    sheet.title = "Scores"

    sheet["A1"] = f"{case_id} - the ten numbers"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (
        "Two numbers per marker, five markers. Percent positive is rounded to the "
        "nearest 5; intensity is on OncoStem's 0-2 scale, banded to one of 0, 0.5, 1, "
        "1.5, 1.75, 2. This is the whole deliverable - there is deliberately no "
        "combined score on this sheet."
    )
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("A2:G2")
    sheet.row_dimensions[2].height = 45

    sheet["A3"] = f"Generated {datetime.now(UTC).isoformat(timespec='seconds')}"
    sheet["A4"] = f"Slides: {case_path}"

    headers = [
        "Marker",
        "Antibody",
        "Compartment",
        "Percent positive (%)",
        "Intensity (0-2)",
        "Intensity band",
        "Tumour cells measured",
        "Status",
    ]
    for column, title in enumerate(headers, start=1):
        sheet.cell(row=6, column=column, value=title)
    _style_header(sheet, row=6)

    deliverable = PatternFill("solid", fgColor=_DELIVERABLE_FILL)
    warn = PatternFill("solid", fgColor=_WARN_FILL)

    row = 7
    for entry in grid.rows:
        spec = panel.spec(entry.marker)
        sheet.cell(row=row, column=1, value=entry.marker)
        sheet.cell(row=row, column=2, value=entry.marker_name)
        sheet.cell(row=row, column=3, value=spec.compartment.value)

        if entry.state == "scored":
            percent = sheet.cell(row=row, column=4, value=entry.percent)
            intensity = sheet.cell(row=row, column=5, value=entry.intensity)
            percent.fill = deliverable
            intensity.fill = deliverable
            percent.font = Font(bold=True)
            intensity.font = Font(bold=True)
            sheet.cell(row=row, column=6, value=entry.intensity_label)
            sheet.cell(row=row, column=7, value=entry.cells)
            sheet.cell(row=row, column=8, value="scored")
        else:
            for column in (4, 5, 6, 7):
                cell = sheet.cell(row=row, column=column, value="not measured")
                cell.fill = warn
            status = sheet.cell(row=row, column=8, value=entry.state)
            status.fill = warn
            sheet.cell(row=row + 0, column=9, value=entry.detail)
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="READ THIS BEFORE USING THE NUMBERS ABOVE")
    sheet.cell(row=row, column=1).font = Font(bold=True)
    row += 1
    for line in _headline_caveats(scores):
        sheet.cell(row=row, column=1, value=line)
        sheet.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        sheet.row_dimensions[row].height = 30
        row += 1

    _autosize(sheet)


def _headline_caveats(scores: dict) -> list[str]:
    """The caveats that apply to the sheet as a whole, deduplicated."""
    lines: list[str] = []
    seen: set[str] = set()
    for result in scores.values():
        for caveat in getattr(result, "caveats", []):
            key = caveat.split(".")[0]
            if key in seen:
                continue
            seen.add(key)
            lines.append(caveat)
    if not lines:
        lines.append("No caveats were raised for the markers on this sheet.")
    lines.append(
        "Full detail is on the Caveats sheet; the arithmetic behind each pair is on "
        "Workings."
    )
    return lines


# --- sheet 2: against the pathologists -------------------------------------


def _agreement_sheet(book, scores: dict, case_id: str) -> None:
    """Our two numbers beside the readers', per marker.

    **The standard is the readers' own spread, not a correct answer.** There is
    no true percent positive for a slide: four trained pathologists reading this
    same block produced four different numbers, and the width of that disagreement
    is the resolution of the measurement. So every row prints the four individual
    reads, their consensus, our number, and the gap - and judges the gap against
    +/-10 absolute points, which is the tolerance OncoStem's own re-review rule
    uses and which no reading in their 120 exceeds.

    A marker with no reader data gets a row saying so. A blank would be
    indistinguishable from agreement.
    """
    from openpyxl.styles import Alignment, Font, PatternFill

    from app.services.validation_service import validation_service

    sheet = book.create_sheet("Agreement")
    sheet["A1"] = "Against the pathologists"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (
        "The target is +/-10 absolute percentage points against the four readers' "
        "consensus, because that is the spread those four reach among themselves. "
        "Landing inside it puts this software inside the human spread; landing outside "
        "it does not, and no amount of agreement elsewhere makes up for it."
    )
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("A2:J2")
    sheet.row_dimensions[2].height = 45

    readers = validation_service.readers()
    from app.services.validation_service import _normalise_case, _same_case

    ours_key = _normalise_case(case_id)

    headers = [
        "Marker",
        "Antibody",
        "Ours (%)",
        "Readers' consensus (%)",
        "Difference (points)",
        "Within +/-10?",
        "Reader 1",
        "Reader 2",
        "Reader 3",
        "Reader 4",
        "Readers' own spread",
    ]
    for column, title in enumerate(headers, start=1):
        sheet.cell(row=4, column=column, value=title)
    _style_header(sheet, row=4)

    good = PatternFill("solid", fgColor=_DELIVERABLE_FILL)
    bad = PatternFill("solid", fgColor=_WARN_FILL)

    row = 5
    for letter in panel.SCORED_MARKERS:
        result = scores.get(letter)
        spec = panel.spec(letter)
        sheet.cell(row=row, column=1, value=letter)
        sheet.cell(row=row, column=2, value=spec.full_name)

        cases = readers.get(letter, {})
        match = next((key for key in cases if _same_case(ours_key, key)), None)
        readings = cases.get(match, []) if match else []

        if result is None:
            sheet.cell(row=row, column=3, value="not measured")
            row += 1
            continue

        sheet.cell(row=row, column=3, value=result.percent)

        if not readings:
            sheet.cell(row=row, column=4, value="no reader data")
            row += 1
            continue

        percents = [reading[0] for reading in readings]
        consensus = sum(percents) / len(percents)
        difference = result.percent - consensus
        within = abs(difference) <= 10.0

        sheet.cell(row=row, column=4, value=round(consensus, 2))
        cell = sheet.cell(row=row, column=5, value=round(difference, 2))
        cell.fill = good if within else bad
        verdict = sheet.cell(row=row, column=6, value="yes" if within else "NO")
        verdict.fill = good if within else bad
        verdict.font = Font(bold=not within)

        for offset, value in enumerate(sorted(percents)):
            sheet.cell(row=row, column=7 + offset, value=value)
        sheet.cell(row=row, column=11, value=round(max(percents) - min(percents), 2))
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Intensity, where the sheet carries it")
    sheet.cell(row=row, column=1).font = Font(bold=True)
    row += 1
    for column, title in enumerate(
        ("Marker", "Ours", "Readers' consensus", "Difference", "Reader values"), start=1
    ):
        sheet.cell(row=row, column=column, value=title)
    _style_header(sheet, row=row)
    row += 1

    for letter in panel.SCORED_MARKERS:
        result = scores.get(letter)
        if result is None:
            continue
        cases = readers.get(letter, {})
        match = next((key for key in cases if _same_case(ours_key, key)), None)
        values = [
            reading[1]
            for reading in (cases.get(match, []) if match else [])
            if reading[1] is not None
        ]
        sheet.cell(row=row, column=1, value=letter)
        sheet.cell(row=row, column=2, value=result.intensity)
        if not values:
            sheet.cell(
                row=row,
                column=3,
                value="no intensity column for this marker in the sheet",
            )
        else:
            consensus = sum(values) / len(values)
            sheet.cell(row=row, column=3, value=round(consensus, 4))
            sheet.cell(row=row, column=4, value=round(result.intensity - consensus, 4))
            sheet.cell(row=row, column=5, value=", ".join(str(v) for v in sorted(values)))
        row += 1

    row += 1
    for line in (
        "Source: " + str(validation_service._sheet_path() or "none"),
        "This transcription carries six of the sheet's eleven columns - the five "
        "percent columns and pan-cadherin's intensity. The other four intensity "
        "columns were not legible in the source and are absent rather than guessed.",
        "Reader values are shown sorted, not attributed. Which of the four "
        "pathologists gave which number is not something this comparison depends on.",
    ):
        sheet.cell(row=row, column=1, value=line)
        sheet.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
        sheet.row_dimensions[row].height = 30
        row += 1

    _autosize(sheet)


# --- sheet 3: the arithmetic ----------------------------------------------


def _workings_sheet(book, scores: dict) -> None:
    from openpyxl.styles import Font

    sheet = book.create_sheet("Workings")
    sheet["A1"] = "How each pair was arrived at"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (
        "Counts, then the formula with the actual numbers substituted in, then the "
        "rounding shown as its own line. Nothing on the Scores sheet is arrived at by "
        "arithmetic that is not written out here."
    )

    row = 4
    for letter in panel.SCORED_MARKERS:
        result = scores.get(letter)
        if result is None:
            continue
        sheet.cell(row=row, column=1, value=f"{letter} - {result.marker_name}")
        sheet.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1

        for column, title in enumerate(("Step", "Expression", "Value"), start=1):
            sheet.cell(row=row, column=column, value=title)
        _style_header(sheet, row=row)
        row += 1

        for step in result.cascade:
            sheet.cell(row=row, column=1, value=step.label)
            sheet.cell(row=row, column=2, value=step.expression)
            sheet.cell(row=row, column=3, value=step.value)
            row += 1

        sheet.cell(row=row, column=1, value="Cut points used")
        sheet.cell(row=row, column=2, value=(
            f"OD cuts {result.od_cuts} (0/1+, 1+/2+, 2+/3+);  "
            f"{result.second_measure.replace('_', ' ')} >= {result.second_min}"
        ))
        sheet.cell(
            row=row,
            column=3,
            value="PROVISIONAL" if result.cuts_provisional else "calibrated",
        )
        row += 1

        sheet.cell(row=row, column=1, value="Cells by bin")
        sheet.cell(row=row, column=2, value=(
            f"0: {result.bin_counts[0]:,}   1+: {result.bin_counts[1]:,}   "
            f"2+: {result.bin_counts[2]:,}   3+: {result.bin_counts[3]:,}"
        ))
        sheet.cell(row=row, column=3, value=f"{result.cells:,} total")
        row += 2

    _autosize(sheet)
    _wrap(sheet, "B", width=80)


# --- sheet 4: what has to be known -----------------------------------------


def _caveats_sheet(book, scores: dict) -> None:
    from openpyxl.styles import Font

    sheet = book.create_sheet("Caveats")
    sheet["A1"] = "What has to be known before these numbers are used"
    sheet["A1"].font = Font(bold=True, size=14)

    for column, title in enumerate(("Marker", "Antibody", "Caveat"), start=1):
        sheet.cell(row=3, column=column, value=title)
    _style_header(sheet, row=3)

    row = 4
    for letter in panel.SCORED_MARKERS:
        result = scores.get(letter)
        if result is None:
            continue
        caveats = result.caveats or ["none recorded"]
        for caveat in caveats:
            sheet.cell(row=row, column=1, value=letter)
            sheet.cell(row=row, column=2, value=result.marker_name)
            sheet.cell(row=row, column=3, value=caveat)
            row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Open questions that move these numbers")
    sheet.cell(row=row, column=1).font = Font(bold=True)
    row += 1
    for question in _open_questions(scores):
        sheet.cell(row=row, column=3, value=question)
        row += 1

    _autosize(sheet)
    _wrap(sheet, "C", width=110)


def _open_questions(scores: dict) -> list[str]:
    cut_file = cut_points.cut_set()
    questions = [
        "Q1 - how a cell whose compartment is only partly stained enters the percentage. "
        "The reading in use is '"
        + cut_file.partial_membrane_rule
        + "'. Each marker's row on the Reference sheet shows what the other readings "
        "would have reported.",
        "Q2 - which intensity scale is authoritative. " + cut_file.band_conflict,
        "Q3 - how the invasive regions are combined. Area-weighted is reported; the "
        "plain mean is on the Regions sheet beside it.",
        "Q6 - whether N-cadherin and pan-cadherin are genuinely near-constant at 80 %, "
        "or 80 % is a working convention. Until that is answered, agreement on those two "
        "markers demonstrates nothing.",
    ]
    if any(getattr(result, "cuts_provisional", False) for result in scores.values()):
        questions.insert(
            0,
            "The cut points are PROVISIONAL. They have not been fitted against the 120 "
            "pathologist readings in 6Slide Reports2.xlsx, because that sheet is not on "
            "disk. Both numbers move with them.",
        )
    return questions


# --- sheet 5: the standard vocabulary ---------------------------------------


def _reference_sheet(book, scores: dict) -> None:
    from openpyxl.styles import Font

    sheet = book.create_sheet("Reference")
    sheet["A1"] = "Reference scores - NOT the deliverable"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (
        "The field's standard vocabulary, computed so the pair on the Scores sheet is "
        "legible against it. OncoStem's own sheet has ten columns and no combined score "
        "anywhere, and collapsing the pair into one number throws away information their "
        "model uses. None of the five antibodies in this panel is HER2; the ASCO/CAP "
        "column applies that guideline's vocabulary to a membrane marker's completeness "
        "and is not a HER2 result."
    )
    sheet.merge_cells("A2:J2")
    sheet.row_dimensions[2].height = 60

    headers = [
        "Marker",
        "Antibody",
        "H-score (0-300)",
        "Allred proportion (0-5)",
        "Allred intensity (0-3)",
        "Allred total (0-8)",
        "ASCO/CAP-style call",
        "Percent if partial counted",
        "Percent if partial halved",
        "Percent if partial excluded",
    ]
    for column, title in enumerate(headers, start=1):
        sheet.cell(row=4, column=column, value=title)
    _style_header(sheet, row=4)

    row = 5
    for letter in panel.SCORED_MARKERS:
        result = scores.get(letter)
        if result is None:
            continue
        by_rule = result.percent_by_partial_rule or {}
        values = [
            letter,
            result.marker_name,
            result.h_score,
            result.allred_proportion,
            result.allred_intensity,
            result.allred_total,
            result.her2_call or "not applicable (cytoplasmic)",
            by_rule.get("count"),
            by_rule.get("half"),
            by_rule.get("exclude"),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value)
        row += 1

    _autosize(sheet)


# --- sheet 6: the regions ---------------------------------------------------


def _regions_sheet(book, scores: dict, measurements: dict) -> None:
    from openpyxl.styles import Font

    sheet = book.create_sheet("Regions")
    sheet["A1"] = "Each invasive region's own answer"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (
        "A single percentage is an average over a very non-uniform field. These are the "
        "regions it averages, and the two ways of combining them - which is open "
        "question Q3."
    )

    headers = [
        "Marker",
        "Region",
        "Area (mm2)",
        "Cells",
        "Positive cells",
        "Percent (raw)",
        "Intensity (raw OD)",
    ]
    for column, title in enumerate(headers, start=1):
        sheet.cell(row=4, column=column, value=title)
    _style_header(sheet, row=4)

    row = 5
    for letter in panel.SCORED_MARKERS:
        result = scores.get(letter)
        if result is None:
            continue
        for region in result.regions:
            values = [
                letter,
                region.rank,
                region.area_mm2,
                region.cells,
                region.positive_cells,
                region.percent_raw,
                region.intensity_raw,
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row, column=column, value=value)
            row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Combining the regions")
    sheet.cell(row=row, column=1).font = Font(bold=True)
    row += 1
    for column, title in enumerate(
        ("Marker", "Area-weighted (%)", "Plain mean (%)", "Gap (points)", "Reported"),
        start=1,
    ):
        sheet.cell(row=row, column=column, value=title)
    _style_header(sheet, row=row)
    row += 1
    for letter in panel.SCORED_MARKERS:
        result = scores.get(letter)
        if result is None:
            continue
        values = [
            letter,
            result.percent_area_weighted,
            result.percent_plain_mean,
            result.averaging_gap_points,
            "area-weighted",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value)
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Heterogeneity - percent positive per sampled field")
    sheet.cell(row=row, column=1).font = Font(bold=True)
    row += 1
    for column, title in enumerate(
        ("Marker", "Region", "Field", "x", "y", "Cells", "Percent"), start=1
    ):
        sheet.cell(row=row, column=column, value=title)
    _style_header(sheet, row=row)
    row += 1
    for letter in panel.SCORED_MARKERS:
        result = scores.get(letter)
        if result is None:
            continue
        for tile in result.heterogeneity:
            values = [
                letter,
                tile.region_rank,
                tile.field_index,
                tile.x,
                tile.y,
                tile.cells,
                tile.percent,
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row, column=column, value=value)
            row += 1

    _autosize(sheet)


# --- sheet 7: the run itself ------------------------------------------------


def _run_sheet(book, runs: list[dict], case_id: str, case_path: str) -> None:
    from openpyxl.styles import Font

    sheet = book.create_sheet("Run")
    sheet["A1"] = "How these numbers were produced"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = f"Case {case_id}"
    sheet["A3"] = f"Slides {case_path}"
    sheet["A4"] = f"Written {datetime.now(UTC).isoformat(timespec='seconds')}"
    sheet["A5"] = f"Cut points {cut_points.cut_set().source_path}"

    headers = [
        "Marker",
        "Antibody",
        "H&E upload id",
        "IHC upload id",
        "State",
        "Seconds",
        "Stages",
        "Error",
    ]
    for column, title in enumerate(headers, start=1):
        sheet.cell(row=7, column=column, value=title)
    _style_header(sheet, row=7)

    row = 8
    for run in runs:
        values = [
            run.get("marker"),
            run.get("marker_name"),
            run.get("he_upload_id"),
            run.get("ihc_upload_id"),
            run.get("state"),
            run.get("seconds"),
            json.dumps(run.get("stages", {})),
            run.get("error", ""),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value)
        row += 1

    _autosize(sheet)
    _wrap(sheet, "G", width=60)
    _wrap(sheet, "H", width=60)


__all__ = ["write_case_workbook"]
